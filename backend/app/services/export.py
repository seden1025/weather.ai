"""설문(예측) 응답을 엑셀 파일(.xlsx)로 만드는 로직.

DB나 웹 프레임워크에 의존하지 않고, 이미 계산된 행(dict) 목록만 받아서 파일 바이트를 돌려준다.
"""

from datetime import date, datetime
from io import BytesIO
from statistics import mean

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# (행 dict의 키, 엑셀 머리글, 열 너비)
COLUMNS: list[tuple[str, str, int]] = [
    ("id", "번호", 7),
    ("created_at", "제출 시각(KST)", 19),
    ("predict_date", "예측 날짜", 12),
    ("station_id", "관측지점", 9),
    ("account_email", "로그인 계정", 26),
    ("observed_sky", "관찰: 하늘", 11),
    ("observed_wind", "관찰: 바람", 11),
    ("user_sky", "내 예측: 날씨", 13),
    ("user_temp", "내 예측: 최고기온(°C)", 16),
    ("basis", "예측 근거(선택)", 34),
    ("basis_detail", "예측 근거(직접 서술)", 40),
    ("memo", "메모", 30),
    ("ai_sky", "AI 예측: 날씨", 13),
    ("ai_temp", "AI 예측: 최고기온(°C)", 16),
    ("actual_temp", "실제 최고기온(°C)", 15),
    ("actual_sky", "실제 날씨", 10),
    ("user_err", "내 기온 오차(°C)", 14),
    ("ai_err", "AI 기온 오차(°C)", 14),
    ("temp_winner", "기온 승자", 10),
    ("user_sky_ok", "내 날씨 적중", 12),
    ("ai_sky_ok", "AI 날씨 적중", 12),
    ("reward_contact", "보상 연락처", 20),
    ("contact_consent", "연락처 동의", 11),
]

PENDING = "집계 전"


def sky_bucket(condition: str | None) -> str | None:
    """구름 관측이 없어 판정은 맑음/비/눈 3가지로만 한다. (프론트 predict 페이지와 같은 규칙)"""
    if not condition:
        return None
    return condition if condition in ("비", "눈") else "맑음"


def actual_sky(rained: bool | None, actual_temp: float | None) -> str | None:
    if rained is None:
        return None
    if not rained:
        return "맑음"
    return "눈" if actual_temp is not None and actual_temp <= 1 else "비"


def enrich(row: dict) -> dict:
    """원본 응답에 실제값과 비교한 파생 값(오차, 승자, 적중 여부)을 덧붙인다."""
    out = dict(row)
    actual = row.get("actual_temp")
    user_err = abs(row["user_temp"] - actual) if actual is not None else None
    ai_err = (
        abs(row["ai_temp"] - actual)
        if actual is not None and row.get("ai_temp") is not None
        else None
    )
    out["user_err"] = round(user_err, 1) if user_err is not None else None
    out["ai_err"] = round(ai_err, 1) if ai_err is not None else None
    if user_err is not None and ai_err is not None:
        out["temp_winner"] = "사람" if user_err < ai_err else "AI" if user_err > ai_err else "무승부"
    else:
        out["temp_winner"] = None

    actual_bucket = row.get("actual_sky")
    for key, src in (("user_sky_ok", "user_sky"), ("ai_sky_ok", "ai_sky")):
        mine = sky_bucket(row.get(src))
        out[key] = (mine == actual_bucket) if actual_bucket is not None and mine is not None else None
    return out


def _set(ws, r: int, c: int, value) -> None:
    cell = ws.cell(row=r, column=c)
    if isinstance(value, bool):
        cell.value = "O" if value else "X"
    elif isinstance(value, (int, float, datetime, date)):
        cell.value = value
        if isinstance(value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm"
        elif isinstance(value, date):
            cell.number_format = "yyyy-mm-dd"
    elif value is None:
        cell.value = None
    else:
        # 참여자가 입력한 글이 '=', '+', '-', '@'로 시작해도 엑셀이 수식으로 실행하지 않도록
        # 반드시 문자열 셀로 저장한다. (전화번호의 앞자리 0도 함께 보존된다)
        cell.value = str(value)
        cell.data_type = "s"
        cell.number_format = "@"


def _summary_rows(rows: list[dict]) -> list[tuple[str, object]]:
    n = len(rows)
    judged = [r for r in rows if r["temp_winner"] is not None]
    user_errs = [r["user_err"] for r in rows if r["user_err"] is not None]
    ai_errs = [r["ai_err"] for r in rows if r["ai_err"] is not None]

    def count(key: str, val) -> int:
        return sum(1 for r in rows if r[key] == val)

    def rate(ok_key: str) -> str:
        done = [r[ok_key] for r in rows if r[ok_key] is not None]
        return f"{sum(done)}/{len(done)}" if done else "-"

    return [
        ("총 응답 수", n),
        ("로그인 상태로 제출", sum(1 for r in rows if r.get("account_email"))),
        ("동의하고 연락처를 남김", sum(1 for r in rows if r.get("reward_contact"))),
        ("실제 결과가 집계된 응답", len(judged)),
        ("기온 승자: 사람", count("temp_winner", "사람")),
        ("기온 승자: AI", count("temp_winner", "AI")),
        ("기온 승자: 무승부", count("temp_winner", "무승부")),
        ("내 평균 기온 오차(°C)", round(mean(user_errs), 2) if user_errs else "-"),
        ("AI 평균 기온 오차(°C)", round(mean(ai_errs), 2) if ai_errs else "-"),
        ("내 날씨 적중(맑음/비/눈 기준)", rate("user_sky_ok")),
        ("AI 날씨 적중(맑음/비/눈 기준)", rate("ai_sky_ok")),
    ]


def build_workbook(raw_rows: list[dict]) -> bytes:
    rows = [enrich(r) for r in raw_rows]

    wb = Workbook()
    ws = wb.active
    ws.title = "응답"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    for c, (_, title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 32

    for r, row in enumerate(rows, start=2):
        for c, (key, _, _) in enumerate(COLUMNS, start=1):
            value = row.get(key)
            if key in ("actual_temp", "actual_sky") and value is None:
                value = PENDING
            _set(ws, r, c, value)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 2)}"

    summary = wb.create_sheet("요약")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 16
    summary.cell(row=1, column=1, value="항목").font = Font(bold=True)
    summary.cell(row=1, column=2, value="값").font = Font(bold=True)
    for r, (label, value) in enumerate(_summary_rows(rows), start=2):
        _set(summary, r, 1, label)
        _set(summary, r, 2, value)
        summary.cell(row=r, column=2).alignment = Alignment(horizontal="right")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
